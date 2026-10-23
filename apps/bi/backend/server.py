"""Local single-user BI server. Access is intentionally bound to loopback only."""
import csv
import copy
import functools
import threading
import hmac
import io
import json
import os
import re
import sqlite3
from datetime import datetime,date
from pathlib import Path
from uuid import uuid4
from typing import Optional
from contextlib import asynccontextmanager

from fastapi import FastAPI, File, UploadFile, HTTPException, Body, Request, Form
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from openpyxl import load_workbook
from .semantic import Engine, SemanticError, columns, model_validate, validate_relation, parse_dax
from .filtering import typed, blank
from .cleaning import prepare, fingerprint, related_orphans
from .workspaces import WorkspaceLibrary, empty_project
from .suite_live import append_batch, SCHEMA as SUITE_INGEST_SCHEMA, TABLE as SUITE_TABLE

ROOT=Path(__file__).resolve().parents[1]
FRONT=ROOT/'frontend'
STATE=Path(os.environ.get('ABBI_DATA_DIR',str(ROOT/'data')))/'workspace.json'
MAX_UPLOAD=18*1024*1024
WRITE_LOCK=threading.RLock()

def serialized(fn):
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        with WRITE_LOCK:
            return fn(*args, **kwargs)
    return wrapper



def demo():
    import random
    rng=random.Random(1120)
    equips=[{'EquipoID':1,'Linea':'Línea 01','Area':'Mezclas','Familia':'Mezclador'},
            {'EquipoID':2,'Linea':'Línea 02','Area':'Mezclas','Familia':'Mezclador'},
            {'EquipoID':3,'Linea':'Línea 03','Area':'Extrusión','Familia':'Extrusora'},
            {'EquipoID':4,'Linea':'Línea 04','Area':'Extrusión','Familia':'Extrusora'}]
    turnos=[{'TurnoID':1,'Turno':'Mañana'},{'TurnoID':2,'Turno':'Tarde'},{'TurnoID':3,'Turno':'Noche'}]
    dates=[];lecturas=[]
    from datetime import timedelta
    for d in range(42):
        dia=date(2026,9,1)+timedelta(days=d)
        dates.append({'Fecha':dia.isoformat(),'Mes':dia.strftime('%Y-%m'),'Semana':f'S{dia.isocalendar().week}','DiaSemana':dia.weekday()+1})
        for e in equips:
            for turno in turnos:
                produced=round(rng.uniform(580,1180)*(1.0 if e['Area']=='Mezclas' else 0.88),1)
                energy=round(produced*rng.uniform(.19,.31),2)
                downtime=round(rng.uniform(0,32),1)
                reject=round(produced*rng.uniform(.004,.04),2)
                lecturas.append({'Fecha':dia.isoformat(),'EquipoID':e['EquipoID'],'TurnoID':turno['TurnoID'],
                    'Produccion_kg':produced,'Energia_kWh':energy,'Parada_min':downtime,
                    'Rechazo_kg':reject,'Temp_C':round(rng.uniform(62,92),2),
                    'Presion_bar':round(rng.uniform(3.4,6.9),2)})
    model={'format':'AB-BI Project','version':3,'name':'Demo industrial · 4 líneas',
           'tables':{'Lecturas':lecturas,'Equipos':equips,'Turnos':turnos,'Calendario':dates},
           'relationships':[
             {'from_table':'Lecturas','from_column':'EquipoID','to_table':'Equipos','to_column':'EquipoID','active':True},
             {'from_table':'Lecturas','from_column':'TurnoID','to_table':'Turnos','to_column':'TurnoID','active':True},
             {'from_table':'Lecturas','from_column':'Fecha','to_table':'Calendario','to_column':'Fecha','active':True}],
           'measures':{
             'Producción total':'SUM(Lecturas[Produccion_kg])',
             'Energía total':'SUM(Lecturas[Energia_kWh])',
             'Consumo específico':'DIVIDE([Energía total], [Producción total])',
             'Temperatura media':'AVERAGE(Lecturas[Temp_C])',
             'Tiempo parada':'SUM(Lecturas[Parada_min])',
             'Lotes':'COUNTROWS(Lecturas)',
             'Rechazos':'SUM(Lecturas[Rechazo_kg])',
             'Rendimiento útil %':'DIVIDE([Producción total] - [Rechazos], [Producción total]) * 100',
             'Producción mezclas':'CALCULATE([Producción total], Equipos[Area] = "Mezclas")',
             '% del total por línea':'DIVIDE([Producción total], CALCULATE([Producción total], REMOVEFILTERS(Equipos[Linea]))) * 100',
             'Producción ponderada':'SUMX(Lecturas, Lecturas[Produccion_kg] * 0.98)',
           },
           'visuals':[
             {'id':'k1','type':'kpi','title':'Producción total','measure':'Producción total'},
             {'id':'k2','type':'kpi','title':'Energía consumida','measure':'Energía total'},
             {'id':'k3','type':'kpi','title':'Consumo específico','measure':'Consumo específico'},
             {'id':'k4','type':'kpi','title':'Rendimiento útil (%)','measure':'Rendimiento útil %'},
             {'id':'v1','type':'bar','title':'Producción por línea','measure':'Producción total','dimension':'Equipos[Linea]','sort':'desc','limit':12},
             {'id':'v2','type':'line','title':'Producción por fecha','measure':'Producción total','dimension':'Calendario[Fecha]','sort':'label','limit':40},
             {'id':'v3','type':'donut','title':'Distribución por área','measure':'Producción total','dimension':'Equipos[Area]','sort':'desc','limit':12},
             {'id':'v4','type':'hbar','title':'Paradas por turno','measure':'Tiempo parada','dimension':'Turnos[Turno]','sort':'desc','limit':12},
           ],
           'filters':{}}
    return model


LIBRARY=WorkspaceLibrary(STATE.parent,model_validate)

def hydrate():
    return LIBRARY.initialize()

MODEL=hydrate()
STATE=LIBRARY.state_path


def persist():
    STATE.parent.mkdir(parents=True,exist_ok=True)
    temp=STATE.with_suffix('.tmp')
    temp.write_text(json.dumps(MODEL,ensure_ascii=False,indent=1,allow_nan=False),encoding='utf-8')
    temp.replace(STATE)

def fail(e):
    if isinstance(e,HTTPException):raise e
    if isinstance(e,(SemanticError,ValueError,KeyError,TypeError)):
        raise HTTPException(422,detail=str(e)) from e
    raise e

app=FastAPI(title='AB-BI Studio API',version='0.4.0',docs_url='/api/docs',redoc_url=None)

@app.middleware('http')
async def active_workspace_guard(request:Request,call_next):
    supplied=request.headers.get('x-abbi-workspace')
    allowed=(request.method=='GET' and request.url.path in ('/api/model','/api/workspaces')) or request.url.path.endswith('/activate')
    if supplied and supplied!=LIBRARY.active and not allowed:
        return JSONResponse({'detail':'El espacio activo cambió en otra pestaña. Abre de nuevo tu espacio antes de guardar.'},status_code=409)
    return await call_next(request)

@app.get('/')
def index():return FileResponse(FRONT/'index.html')

@app.get('/classic')
def classic():return FileResponse(FRONT/'legacy_v02.html')

@app.get('/api/health')
def health():return {'ok':True,'version':'0.4.0','engine':'Semantic DAX subset + SQLite read-only SQL'}

def summary():
    return {'workspace_id':LIBRARY.active,'name':MODEL.get('name','Sin título'),'format':MODEL['format'],'version':3,
            'tables':[{'name':n,'rows':len(r),'columns':columns(MODEL,n)} for n,r in MODEL['tables'].items()],
            'relationships':MODEL['relationships'], 'measures':MODEL['measures'],
            'visuals':MODEL['visuals'],'filters':MODEL.get('filters',{}),
            'cleaning_history':MODEL.get('cleaning_history',[])[-20:]}

@app.get('/api/model')
@serialized
def get_model():return summary()

@app.get('/api/table/{name}')
@serialized
def table(name:str,limit:int=50,offset:int=0):
    if name not in MODEL['tables']:raise HTTPException(404,'Tabla no encontrada')
    rows=MODEL['tables'][name]
    return {'name':name,'total':len(rows),'rows':rows[max(0,offset):max(0,offset)+min(max(limit,1),500)]}

@app.get('/api/profile/{name}')
@serialized
def profile(name:str):
    if name not in MODEL['tables']:raise HTTPException(404,'Tabla inexistente')
    rows=MODEL['tables'][name]
    result=[]
    for c in columns(MODEL,name):
        vals=[r.get(c) for r in rows]
        valid=[v for v in vals if not blank(v)]
        nums=[]
        for v in valid:
            try:
                from .semantic import as_num
                n=as_num(v)
                if n is not None:nums.append(n)
            except Exception:pass
        ct={}
        for v in valid:
            key=str(v);ct[key]=ct.get(key,0)+1
        most=sorted(ct.items(),key=lambda x:-x[1])[:5]
        typ='Número' if valid and len(nums)/len(valid)>.9 else 'Texto/fecha'
        result.append({'name':c,'type':typ,'count':len(valid),'missing':len(vals)-len(valid),
                       'distinct':len(ct),'min':min(nums) if typ=='Número' else None,
                       'max':max(nums) if typ=='Número' else None,
                       'average':sum(nums)/len(nums) if typ=='Número' and nums else None,
                       'most_common':most,
                       'whitespace':sum(isinstance(v,str) and v!=v.strip() for v in vals),
                       'numeric_candidates':len(nums),
                       'non_numeric':len(valid)-len(nums)})
    return {'name':name,'rows':len(rows),'columns':result,
            'duplicate_rows':len(rows)-len({fingerprint(r) for r in rows}),
            'missing_cells':sum(c['missing'] for c in result),
            'whitespace_cells':sum(c['whitespace'] for c in result)}

@app.post('/api/demo')
@serialized
def new_demo():
    global MODEL,STATE
    persist();candidate=demo();LIBRARY.create(candidate);STATE=LIBRARY.state_path;MODEL=candidate;return summary()


def safe_name(source, existing):
    s=re.sub(r'[^\wÁÉÍÓÚÜÑáéíóúüñ .-]','_',str(source).strip())[:70].strip('. ')
    s=s or 'Tabla'
    name=s;i=2
    while name in existing:name=f'{s}_{i}';i+=1
    return name


def row_objects(headers, records):
    headers=[str(h).strip() if h not in (None,'') else f'Campo_{i+1}' for i,h in enumerate(headers)]
    seen=set();unique=[]
    for h in headers:
        name=h;j=2
        while name in seen:name=f'{h}_{j}';j+=1
        seen.add(name);unique.append(name)
    if len(unique)>150:raise SemanticError('Máximo 150 columnas')
    result=[]
    for vals in records:
        if len(result)>=40000: raise SemanticError('Máximo 40.000 filas por tabla')
        result.append({col:(v.isoformat() if isinstance(v,(date,datetime)) else v if isinstance(v,(str,int,float,bool)) or v is None else str(v)) for col,v in zip(unique,vals)})
    return result


def convert_upload(blob:bytes,filename:str):
    ext=Path(filename).suffix.lower()
    imported={}
    if ext=='.xlsx':
        book=load_workbook(io.BytesIO(blob),read_only=True,data_only=True,keep_links=False)
        try:
            for ws in book.worksheets:
                iterator=ws.iter_rows(values_only=True)
                header=next(iterator,None)
                if header is None:continue
                name=safe_name(ws.title,imported)
                imported[name]=row_objects(header,iterator)
        finally:book.close()
    elif ext in ('.csv','.tsv'):
        text=blob.decode('utf-8-sig',errors='replace')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=';,\t|')
        except csv.Error:dialect=csv.excel_tab if ext=='.tsv' else csv.excel
        iterator=csv.reader(io.StringIO(text),dialect)
        header=next(iterator,None)
        if header is None:raise SemanticError('Archivo CSV vacío')
        imported[safe_name(Path(filename).stem,imported)]=row_objects(header,iterator)
    elif ext=='.json':
        obj=json.loads(blob)
        if isinstance(obj,dict) and 'tables' in obj:
            for n,records in obj['tables'].items():
                if not isinstance(records,list) or not all(isinstance(r,dict) for r in records):
                    raise SemanticError('JSON tables debe contener listas de objetos')
                imported[safe_name(n,imported)]=records
        else:
            if isinstance(obj,dict) and isinstance(obj.get('rows'),list):obj=obj['rows']
            if not isinstance(obj,list) or not all(isinstance(r,dict) for r in obj):
                raise SemanticError('JSON debe ser una lista de objetos')
            imported[safe_name(Path(filename).stem,imported)]=obj
    else:raise SemanticError('Formato admitido: CSV, TSV, JSON, XLSX')
    if not imported:raise SemanticError('No se detectaron hojas con datos')
    return imported

@app.post('/api/import')
async def import_file(file:UploadFile=File(...)):
    try:
        blob=await file.read(MAX_UPLOAD+1)
        if len(blob)>MAX_UPLOAD:raise SemanticError('El archivo supera los 18 MB')
        tables=convert_upload(blob,file.filename or 'archivo.csv')
        with WRITE_LOCK:
            existing=dict(MODEL['tables']);names=[]
            for name,rows in tables.items():
                name=safe_name(name,existing)
                existing[name]=rows;names.append(name)
            updated={**MODEL,'tables':existing}
            model_validate(updated)
            MODEL['tables']=existing
            persist()
            return {'added':names,'model':summary()}
    except Exception as e:fail(e)


@app.get('/api/workspaces')
@serialized
def list_workspaces():
    return {'active':LIBRARY.active,'spaces':LIBRARY.list()}

@app.post('/api/workspaces')
@serialized
def create_workspace(p:dict=Body(...)):
    global MODEL,STATE
    try:
        name=p.get('name','Mi espacio de datos');template=p.get('template','blank')
        if template not in ('blank','demo','copy'):raise SemanticError('Plantilla de espacio desconocida')
        candidate=empty_project(name)
        if template=='demo':candidate={**demo(),'name':candidate['name']}
        elif template=='copy':candidate={**copy.deepcopy(MODEL),'name':candidate['name']}
        model_validate(candidate);persist();LIBRARY.create(candidate)
        STATE=LIBRARY.state_path;MODEL=candidate;return summary()
    except Exception as error:fail(error)

@app.post('/api/workspaces/{identifier}/activate')
@serialized
def activate_workspace(identifier:str):
    global MODEL,STATE
    try:
        candidate=LIBRARY.read(identifier);persist();LIBRARY.activate(identifier)
        STATE=LIBRARY.state_path;MODEL=candidate;return summary()
    except Exception as error:fail(error)

@app.post('/api/workspaces/rename')
@serialized
def rename_workspace(p:dict=Body(...)):
    try:
        name=empty_project(p.get('name'))['name']
        MODEL['name']=name;persist();return summary()
    except Exception as error:fail(error)

@app.post('/api/data/import')
async def import_data(files:list[UploadFile]=File(...),destination:str=Form('new'),name:str=Form('Mi espacio de datos'),workspace_id:str=Form('')):
    global MODEL,STATE
    try:
        if destination not in ('new','current'):raise SemanticError('Destino de importación desconocido')
        if not 1<=len(files)<=24:raise SemanticError('Importa entre 1 y 24 archivos por operación')
        parsed=[];total=0
        for file in files:
            blob=await file.read(MAX_UPLOAD+1);total+=len(blob)
            if len(blob)>MAX_UPLOAD or total>30*1024*1024:raise SemanticError('Límite de 18 MB por archivo y 30 MB por importación')
            parsed.append(convert_upload(blob,file.filename or 'datos.csv'))
        with WRITE_LOCK:
            if destination=='current' and workspace_id!=LIBRARY.active:
                raise HTTPException(409,'El espacio activo cambió en otra pestaña. Recarga antes de importar.')
            candidate=empty_project(name) if destination=='new' else copy.deepcopy(MODEL)
            added=[]
            for tables in parsed:
                for table,rows in tables.items():
                    table=safe_name(table,candidate['tables']);candidate['tables'][table]=rows;added.append(table)
                    measure=safe_name('Filas · '+table,candidate['measures'])
                    candidate['measures'][measure]="COUNTROWS('"+table.replace("'","''")+"')"
            model_validate(candidate)
            if destination=='new':
                persist();LIBRARY.create(candidate);STATE=LIBRARY.state_path;MODEL=candidate
            else:
                previous=MODEL;MODEL=candidate
                try:persist()
                except Exception:MODEL=previous;raise
            return {'added':added,'model':summary()}
    except Exception as error:fail(error)

@app.post('/api/relationship')
@serialized
def relationship(rel:dict=Body(...)):
    try:
        if any(all(r.get(k)==rel.get(k) for k in ('from_table','from_column','to_table','to_column')) for r in MODEL['relationships']):
            raise SemanticError('La relación ya existe')
        new=dict(rel);new['active']=True
        validate_relation(MODEL,new)
        MODEL['relationships'].append(new);persist();return summary()
    except Exception as e:fail(e)

@app.delete('/api/relationship/{idx}')
@serialized
def delete_relationship(idx:int):
    if idx<0 or idx>=len(MODEL['relationships']):raise HTTPException(404,'Relación no encontrada')
    MODEL['relationships'].pop(idx);persist();return summary()

@app.put('/api/measure')
@serialized
def put_measure(p:dict=Body(...)):
    try:
        name=str(p.get('name','')).strip();expr=str(p.get('expression','')).strip()
        if not name or len(name)>80 or not expr:raise SemanticError('Indica nombre y expresión')
        ast=parse_dax(expr)
        candidate={**MODEL,'measures':{**MODEL['measures'],name:expr}}
        # Detect unknown functions / tables / recursive dependencies at evaluation time.
        Engine(candidate).measure(name)
        MODEL['measures'][name]=expr;persist();return summary()
    except Exception as e:fail(e)

@app.delete('/api/measure/{name}')
@serialized
def delete_measure(name:str):
    if name not in MODEL['measures']:raise HTTPException(404,'Medida no encontrada')
    if any(v.get('measure')==name for v in MODEL['visuals']):
        raise HTTPException(422,'La medida se utiliza en un gráfico; elimina el gráfico primero')
    MODEL['measures'].pop(name);persist();return summary()

@app.post('/api/evaluate')
@serialized
def evaluate(p:dict=Body(...)):
    try:
        expr=p.get('expression','');ctx=p.get('filters',{})
        if not isinstance(ctx,dict):raise SemanticError('Filtros inválidos')
        return {'result':Engine(MODEL).eval(expr,ctx)}
    except Exception as e:fail(e)

@app.post('/api/dashboard')
@serialized
def dashboard(p:dict=Body(default={})):
    try:
        ctx=p.get('filters',MODEL.get('filters',{}))
        if not isinstance(ctx,dict) or len(ctx)>35:raise SemanticError('Filtros inválidos')
        engine=Engine(MODEL)
        visuals=[]
        for v in MODEL['visuals']:
            try:visuals.append({**v,'result':engine.visualize(v,ctx)})
            except SemanticError as e:visuals.append({**v,'error':str(e)})
        return {'visuals':visuals,'filters':ctx}
    except Exception as e:fail(e)

@app.post('/api/visual')
@serialized
def put_visual(p:dict=Body(...)):
    try:
        name=str(p.get('title','')).strip()[:95]
        typ=p.get('type')
        if not name or typ not in ('bar','hbar','line','area','donut','kpi','table'):
            raise SemanticError('Título o tipo de visualización inválido')
        if p.get('measure') not in MODEL['measures']:raise SemanticError('Selecciona una medida existente')
        visual={'id':str(p.get('id') or uuid4().hex[:10]),'type':typ,'title':name,'measure':p['measure'],
                'dimension':str(p.get('dimension') or ''),'sort':p.get('sort','desc'),'limit':min(40,max(2,int(p.get('limit',12))))}
        Engine(MODEL).visualize(visual,MODEL.get('filters',{}))
        pos=next((i for i,v in enumerate(MODEL['visuals']) if v['id']==visual['id']),-1)
        if pos>=0:MODEL['visuals'][pos]=visual
        else:
            if len(MODEL['visuals'])>=32:raise SemanticError('Máximo 32 visualizaciones')
            MODEL['visuals'].append(visual)
        persist();return summary()
    except Exception as e:fail(e)

@app.delete('/api/visual/{visual_id}')
@serialized
def del_visual(visual_id:str):
    MODEL['visuals']=[v for v in MODEL['visuals'] if v['id']!=visual_id]
    persist();return summary()

@app.post('/api/filters')
@serialized
def save_filters(filters:dict=Body(...)):
    try:
        if len(filters)>25:raise SemanticError('Máximo 25 filtros')
        engine=Engine(MODEL)
        for col in filters:engine.distinct(col,{},1)
        filters=engine.context(filters)
        MODEL['filters']=filters;persist();return summary()
    except Exception as e:fail(e)

@app.get('/api/distinct')
@serialized
def distinct(field:str,search:str=''):
    try:
        vals=Engine(MODEL).distinct(field,{},40001)
        if search:vals=[v for v in vals if search.casefold() in str(v).casefold()]
        return {'values':vals[:300],'total':len(vals),'truncated':len(vals)>300}
    except Exception as e:fail(e)

@app.get('/api/export')
@serialized
def export_project():
    payload=json.dumps(MODEL,ensure_ascii=False,indent=2,allow_nan=False)
    return JSONResponse(MODEL,headers={'Content-Disposition':'attachment; filename="AB_BI_proyecto_v04.abbi"'})

@app.post('/api/project/import')
async def import_project(file:UploadFile=File(...)):
    global MODEL,STATE
    try:
        blob=await file.read(30*1024*1024+1)
        if len(blob)>30*1024*1024:raise SemanticError('El proyecto supera 30 MB')
        model=json.loads(blob)
        if model.get('format')=='AB-BI Studio Project' and model.get('version') in (1,2):
            rows=model.get('dataset',{}).get('rows',[])
            name=safe_name(model.get('dataset',{}).get('name','Importado_v02'),{})
            model={'format':'AB-BI Project','version':3,'name':name,'tables':{name:rows},'relationships':[],
                   'measures':{'Recuento':f'COUNTROWS({name})'},'visuals':[], 'filters':{}}
        if model.get('format')!='AB-BI Project' or model.get('version')!=3:
            raise SemanticError('Formato de proyecto incompatible')
        model_validate(model)
        model['filters']=Engine(model).context(model.get('filters',{}))
        with WRITE_LOCK:
            LIBRARY.create(model);STATE=LIBRARY.state_path;MODEL=model;return summary()
    except Exception as e:fail(e)

# A separate read-only SQL exploration engine, built from tables in the semantic model.
# SQLite is available on every CPython; DuckDB can be added in later builds.
@app.post('/api/sql')
@serialized
def run_sql(p:dict=Body(...)):
    query=str(p.get('query','')).strip()
    if len(query)>5000 or not re.match(r'^\s*(SELECT|WITH)\b',query,re.I):
        raise HTTPException(422,'Solo se aceptan consultas SELECT o WITH SELECT')
    db=sqlite3.connect(':memory:')
    try:
        for table,rows in MODEL['tables'].items():
            cols=columns(MODEL,table)
            if not cols:continue
            q=lambda x:'"'+x.replace('"','""')+'"'
            spec=[]
            for col in cols:
                values=[r.get(col) for r in rows[:300]]
                ty='REAL' if values and sum(isinstance(x,(int,float)) and not isinstance(x,bool) for x in values)>len(values)*.75 else 'TEXT'
                spec.append(f'{q(col)} {ty}')
            db.execute(f'CREATE TABLE {q(table)} ({", ".join(spec)})')
            placeholders=','.join('?' for _ in cols)
            db.executemany(f'INSERT INTO {q(table)} VALUES ({placeholders})',
                [[str(r.get(c)) if isinstance(r.get(c),(date,datetime)) else r.get(c) for c in cols] for r in rows])
        db.commit()
        db.execute('PRAGMA query_only = ON')
        ticks=[0]
        def progress():
            ticks[0]+=1
            return 1 if ticks[0]>20000 else 0
        db.set_progress_handler(progress,1000)
        cursor=db.execute(query)
        columns_out=[v[0] for v in cursor.description or []]
        data=cursor.fetchmany(501)
        return {'columns':columns_out,'rows':[list(r) for r in data[:500]],'truncated':len(data)>500,'engine':'SQLite'}
    except sqlite3.Error as e:
        raise HTTPException(422,str(e)) from e
    finally:db.close()

@app.post('/api/table-query')
@serialized
def query_table(p:dict=Body(...)):
    try:
        name=p.get('table')
        if name not in MODEL['tables']:raise SemanticError('Tabla no encontrada')
        revision=fingerprint(MODEL)
        if p.get('revision') and p['revision']!=revision:
            raise HTTPException(409,'El proyecto cambió durante la exportación. Vuelve a exportar.')
        engine=Engine(MODEL)
        rows=engine.filtered_rows(name,engine.context(p.get('filters',{})))
        query=str(p.get('search','')).strip().casefold()
        if len(query)>500:raise SemanticError('Búsqueda demasiado larga')
        if query:rows=[r for r in rows if any(query in str(v).casefold() for v in r.values() if v is not None)]
        col=p.get('sort_column','')
        if col:
            if col not in columns(MODEL,name):raise SemanticError('Columna de orden inexistente')
            kind=p.get('sort_type','text')
            if kind not in ('text','number','date'):raise SemanticError('Tipo de orden no válido')
            def sort_key(row):
                value=row.get(col)
                try:return (0,typed(value,kind)) if not blank(value) else (1,'')
                except (ValueError,TypeError,OverflowError):return (1,str(value))
            rows=sorted(rows,key=sort_key,reverse=p.get('sort_direction')=='desc')
        offset=max(0,int(p.get('offset',0)));limit=min(500,max(1,int(p.get('limit',40))))
        return {'rows':rows[offset:offset+limit],'total':len(rows),'source_total':len(MODEL['tables'][name]),
                'columns':columns(MODEL,name),'revision':revision}
    except Exception as exc:fail(exc)


def undo_record():
    path=STATE.parent/'cleaning_undo.json'
    if not path.exists():return None
    try:
        record=json.loads(path.read_text(encoding='utf-8'))
        if not re.fullmatch(r'[a-f0-9]{32}\.json',record.get('backup','')):return None
        if record.get('revision')!=fingerprint(MODEL):return None
        if not (STATE.parent/'backups'/record['backup']).exists():return None
        return record
    except (OSError,ValueError,TypeError):return None


@app.get('/api/cleaning/status')
@serialized
def cleaning_status():
    record=undo_record()
    return {'can_undo':record is not None,'history':MODEL.get('cleaning_history',[])[-20:],
            'table':record.get('table') if record else None,'orphans':related_orphans(MODEL)}


@app.post('/api/cleaning/preview')
@serialized
def cleaning_preview(p:dict=Body(...)):
    try:
        _,preview=prepare(MODEL,p.get('table'),p.get('steps'))
        return preview
    except Exception as exc:fail(exc)


@app.post('/api/cleaning/apply')
@serialized
def cleaning_apply(p:dict=Body(...)):
    global MODEL
    try:
        if not p.get('token'):raise SemanticError('Genera una vista previa antes de aplicar')
        candidate,preview=prepare(MODEL,p.get('table'),p.get('steps'))
        if p['token']!=preview['token']:
            raise HTTPException(409,'Los datos o los pasos han cambiado. Genera otra vista previa.')
        if not preview['changed_rows'] and not preview['removed_rows']:
            raise SemanticError('La receta no produce cambios')
        previous=MODEL
        folder=STATE.parent/'backups';folder.mkdir(parents=True,exist_ok=True)
        backup=uuid4().hex+'.json'
        (folder/backup).write_text(json.dumps(previous,ensure_ascii=False,allow_nan=False),encoding='utf-8')
        entry={'id':uuid4().hex,'at':datetime.now().astimezone().isoformat(timespec='seconds'),
               'table':p['table'],'steps':copy.deepcopy(p['steps']),
               'before_rows':preview['before_rows'],'after_rows':preview['after_rows'],
               'changed_cells':preview['changed_cells'],'invalid_values':preview['invalid_values']}
        candidate['cleaning_history']=[*candidate.get('cleaning_history',[]),entry][-100:]
        record={'backup':backup,'revision':fingerprint(candidate),'backup_revision':fingerprint(previous),'table':p['table']}
        checkpoint=STATE.parent/'cleaning_undo.json'
        previous_checkpoint=checkpoint.read_bytes() if checkpoint.exists() else None
        staged=checkpoint.with_suffix('.tmp')
        staged.write_text(json.dumps(record),encoding='utf-8')
        staged.replace(checkpoint)
        MODEL=candidate
        try:persist()
        except Exception:
            MODEL=previous
            if previous_checkpoint is None:checkpoint.unlink(missing_ok=True)
            else:checkpoint.write_bytes(previous_checkpoint)
            raise
        return {'model':summary(),'preview':preview,'can_undo':True}
    except Exception as exc:fail(exc)


@app.post('/api/cleaning/undo')
@serialized
def cleaning_undo():
    global MODEL
    try:
        record=undo_record()
        if record is None:raise HTTPException(409,'No hay una limpieza que se pueda deshacer sin sobrescribir cambios posteriores.')
        restored=json.loads((STATE.parent/'backups'/record['backup']).read_text(encoding='utf-8'))
        if fingerprint(restored)!=record.get('backup_revision'):
            raise HTTPException(409,'La copia de seguridad ha cambiado o está dañada. No se ha sustituido el proyecto.')
        model_validate(restored)
        previous=MODEL;MODEL=restored
        try:persist()
        except Exception:
            MODEL=previous
            raise
        (STATE.parent/'cleaning_undo.json').unlink(missing_ok=True)
        return summary()
    except Exception as exc:fail(exc)



# Optional live virtual telemetry ingestion. Explicit opt-in via environment.
def suite_write_authorized(request: Request):
    expected=os.environ.get("ABBI_SUITE_WRITE_TOKEN","")
    if len(expected)<32:
        raise HTTPException(503,"Ingesta Suite desactivada: configura ABBI_SUITE_WRITE_TOKEN")
    supplied=request.headers.get("x-abbi-suite-token","")
    if not hmac.compare_digest(supplied,expected):
        raise HTTPException(403,"Token de ingesta Suite inválido")


@app.get("/api/suite/summary")
@serialized
def suite_public_summary():
    enabled=len(os.environ.get("ABBI_SUITE_WRITE_TOKEN",""))>=32
    state=MODEL.get("suite_live",{})
    if not isinstance(state,dict):state={}
    return {"enabled":enabled,"table":SUITE_TABLE,
            "rows":len(MODEL["tables"].get(SUITE_TABLE,[])) if state.get("table")==SUITE_TABLE else 0,
            "revision":str(state.get("session_id",""))+":"+str(state.get("last_sequence",0))}


@app.get("/api/suite/status")
@serialized
def suite_status(request:Request):
    suite_write_authorized(request)
    state=MODEL.get("suite_live",{})
    if not isinstance(state,dict):state={}
    return {"schema":"msrobotika.bi.suite.status.v1",
            "table":SUITE_TABLE,"session_id":state.get("session_id"),
            "last_sequence":state.get("last_sequence",0),
            "rows":len(MODEL["tables"].get(SUITE_TABLE,[])) if state.get("table")==SUITE_TABLE else 0}


@app.post("/api/suite/ingest")
@serialized
def suite_ingest(request:Request,p:dict=Body(...)):
    global MODEL
    suite_write_authorized(request)
    try:
        if p.get("workspace_id") and p["workspace_id"]!=LIBRARY.active:
            raise HTTPException(409,"La captura se detuvo porque cambió el espacio activo de BI")
        if p.get("schema")!=SUITE_INGEST_SCHEMA:
            raise SemanticError("Versión de ingesta Suite incompatible")
        previous=MODEL
        candidate,added,skipped=append_batch(previous,p.get("rows"),p.get("session_id"))
        if added:
            MODEL=candidate
            try:persist()
            except Exception:
                MODEL=previous
                raise
        state=MODEL.get("suite_live",{})
        return {"inserted":added,"duplicates":skipped,
                "last_sequence":state.get("last_sequence",0),
                "table":SUITE_TABLE}
    except Exception as exc:fail(exc)

app.mount('/static',StaticFiles(directory=FRONT),name='static')

if __name__=='__main__':
    import uvicorn
    uvicorn.run(app,host='127.0.0.1',port=int(os.getenv('AB_BI_PORT','8765')))
